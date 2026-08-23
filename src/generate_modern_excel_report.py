from __future__ import annotations

import csv
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]

REPORT_CSV = ROOT / "data" / "weekly_report.csv"
SNAPSHOTS_CSV = ROOT / "data" / "instagram_snapshots.csv"
HISTORICAL_CSV = ROOT / "data" / "historical.csv"
ERRORS_CSV = ROOT / "data" / "profile_validation_errors.csv"

OUTPUT_XLSX = ROOT / "data" / "NBA_Competitor_Report_Latest.xlsx"

NBA_HANDLE = "nbaaustralia_official"


# ============================================================
# COLOURS
# ============================================================

INK = "0A0D12"
INK_2 = "151A22"
SURFACE = "F6F8FB"
CARD = "FFFFFF"
LINE = "E6EAF0"

TEXT = "111827"
MUTED = "667085"

CYAN = "00B8D9"
GREEN = "12B76A"
RED = "F04438"

AMBER_BG = "FFF4E5"
AMBER_LIGHT = "FFFAEB"
AMBER_TEXT = "B54708"

WHITE = "FFFFFF"


# ============================================================
# HELPERS
# ============================================================

def read_csv(path: Path):
    if not path.exists():
        return []

    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def as_int(value):
    if value is None:
        return None

    value = str(value).strip()

    if value == "":
        return None

    try:
        return int(float(value.replace(",", "")))
    except Exception:
        return None


def as_float(value):
    if value is None:
        return None

    value = str(value).strip()

    if value == "":
        return None

    try:
        return float(value.replace("%", "").replace(",", ""))
    except Exception:
        return None


def parse_date(value):
    if not value:
        return None

    value = str(value).strip()

    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass

    return None


def apply_fill(cell, colour):
    cell.fill = PatternFill("solid", fgColor=colour)


def set_border(cell, colour=LINE):
    side = Side(style="thin", color=colour)

    cell.border = Border(
        left=side,
        right=side,
        top=side,
        bottom=side,
    )


def safe_growth_percent(row):
    """
    weekly_report.csv stores growth_percent_30d_plus
    as percentage points.

    Example:
        Previous followers: 100
        Current followers:  101
        CSV growth value:    1.00

    openpyxl / Excel percentage cells expect a decimal ratio,
    so 1.00% must be stored as 0.01.

    Therefore every valid report percentage is divided by 100.
    """

    value = as_float(
        row.get("growth_percent_30d_plus")
    )

    if value is None:
        return None

    return value / 100


def autosize(ws, minimum=8, maximum=40):
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)

        max_len = 0

        for cell in col_cells:
            value = cell.value

            if value is not None:
                max_len = max(max_len, len(str(value)))

        ws.column_dimensions[letter].width = min(
            maximum,
            max(minimum, max_len + 3),
        )


def style_header_row(ws, row=1):
    for cell in ws[row]:
        apply_fill(cell, INK)
        cell.font = Font(
            name="Aptos",
            size=9,
            bold=True,
            color=WHITE,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws.row_dimensions[row].height = 24


def style_body(ws, start_row=2):
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            cell.font = Font(
                name="Aptos",
                size=10,
                color=TEXT,
            )

            cell.alignment = Alignment(
                vertical="center",
            )

        ws.row_dimensions[cell.row].height = 21


def add_kpi_card(
    ws,
    start_col,
    end_col,
    start_row,
    label,
    value,
    subtitle,
    value_colour=TEXT,
):
    ws.merge_cells(
        start_row=start_row,
        start_column=start_col,
        end_row=start_row,
        end_column=end_col,
    )

    ws.merge_cells(
        start_row=start_row + 1,
        start_column=start_col,
        end_row=start_row + 3,
        end_column=end_col,
    )

    ws.merge_cells(
        start_row=start_row + 4,
        start_column=start_col,
        end_row=start_row + 4,
        end_column=end_col,
    )

    for row in range(start_row, start_row + 5):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)

            apply_fill(cell, CARD)
            set_border(cell)

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    label_cell = ws.cell(start_row, start_col)

    label_cell.value = label
    label_cell.font = Font(
        name="Aptos",
        size=9,
        bold=True,
        color=MUTED,
    )

    value_cell = ws.cell(start_row + 1, start_col)

    value_cell.value = value
    value_cell.font = Font(
        name="Aptos Display",
        size=24,
        bold=True,
        color=value_colour,
    )

    value_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    subtitle_cell = ws.cell(start_row + 4, start_col)

    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(
        name="Aptos",
        size=9,
        color=MUTED,
    )

    subtitle_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


# ============================================================
# LOAD DATA
# ============================================================

report = read_csv(REPORT_CSV)

if not report:
    raise RuntimeError(
        f"No report data found in {REPORT_CSV}"
    )

snapshots = read_csv(SNAPSHOTS_CSV)
historical = read_csv(HISTORICAL_CSV)
errors = read_csv(ERRORS_CSV)


for row in report:
    row["_rank"] = as_int(row.get("rank")) or 9999
    row["_followers"] = as_int(row.get("current_followers"))
    row["_comparison_followers"] = as_int(
        row.get("comparison_followers")
    )
    row["_change"] = as_int(
        row.get("follower_change_30d_plus")
    )
    row["_growth"] = safe_growth_percent(row)
    row["_posts"] = as_int(
        row.get("current_total_posts")
    )
    row["_posts_change"] = as_int(
        row.get("posts_since_previous_snapshot")
    )
    row["_comparison_days"] = as_int(
        row.get("comparison_days")
    )


report = sorted(
    report,
    key=lambda r: r["_rank"],
)


latest_date = report[0].get("date", "")

nba = next(
    (
        row
        for row in report
        if row.get("handle") == NBA_HANDLE
    ),
    None,
)

if nba is None:
    raise RuntimeError(
        f"NBA account @{NBA_HANDLE} missing from weekly report."
    )


nba_followers = nba["_followers"] or 0
nba_rank = nba["_rank"]
nba_change = nba["_change"]
nba_growth = nba["_growth"]
nba_days = nba["_comparison_days"]


accounts_ahead = sum(
    1
    for row in report
    if str(row.get("ahead_of_nba", "")).strip().lower()
    in ("true", "yes", "1")
)


# ============================================================
# WORKBOOK
# ============================================================

wb = Workbook()

report_ws = wb.active
report_ws.title = "Report"

ranking_ws = wb.create_sheet("Competitor Ranking")
growth_ws = wb.create_sheet("30-Day Growth")
activity_ws = wb.create_sheet("Posting Activity")
issues_ws = wb.create_sheet("Issues")
history_ws = wb.create_sheet("History")


# ============================================================
# REPORT DASHBOARD
# ============================================================

report_ws.sheet_view.showGridLines = False

for row in report_ws.iter_rows(
    min_row=1,
    max_row=45,
    min_col=1,
    max_col=12,
):
    for cell in row:
        apply_fill(cell, SURFACE)


report_ws.merge_cells("A1:L3")

title = report_ws["A1"]
title.value = "NBA COMPETITOR REPORT"

apply_fill(title, INK)

title.font = Font(
    name="Aptos Display",
    size=22,
    bold=True,
    color=WHITE,
)

title.alignment = Alignment(
    horizontal="left",
    vertical="center",
)


for row in report_ws["A1:L3"]:
    for cell in row:
        apply_fill(cell, INK)


report_ws.merge_cells("A4:L4")

subtitle = report_ws["A4"]

pretty_date = latest_date

parsed_latest = parse_date(latest_date)

if parsed_latest:
    pretty_date = parsed_latest.strftime("%d %B %Y")


subtitle.value = (
    f"Instagram competitor growth • "
    f"{pretty_date} • 30-day+ comparison"
)

subtitle.font = Font(
    name="Aptos",
    size=10,
    color="AAB2C0",
)

subtitle.alignment = Alignment(
    vertical="center",
)

for row in report_ws["A4:L4"]:
    for cell in row:
        apply_fill(cell, INK)


for cell in report_ws["A5:L5"][0]:
    apply_fill(cell, CYAN)

report_ws.row_dimensions[5].height = 4


growth_text = "—"

if nba_growth is not None:
    growth_text = f"{nba_growth:+.2%}"


growth_sub = "No 30-day comparison yet"

if nba_change is not None:
    growth_sub = (
        f"{nba_change:+,} followers"
    )

    if nba_days:
        growth_sub += f" • {nba_days} days"


ahead_sub = (
    "No accounts ahead"
    if accounts_ahead == 0
    else (
        "1 account ahead"
        if accounts_ahead == 1
        else f"{accounts_ahead} accounts ahead"
    )
)


add_kpi_card(
    report_ws,
    1,
    3,
    7,
    "NBA FOLLOWERS",
    f"{nba_followers:,}",
    "Exact current count",
)

add_kpi_card(
    report_ws,
    4,
    6,
    7,
    "FOLLOWER RANK",
    f"#{nba_rank}",
    f"of {len(report)} tracked accounts",
)

add_kpi_card(
    report_ws,
    7,
    9,
    7,
    "30-DAY+ GROWTH",
    growth_text,
    growth_sub,
    GREEN if (nba_growth or 0) >= 0 else RED,
)

add_kpi_card(
    report_ws,
    10,
    12,
    7,
    "ACCOUNTS AHEAD",
    accounts_ahead,
    ahead_sub,
)


# ============================================================
# TOP FOLLOWERS
# ============================================================

report_ws.merge_cells("A14:F14")

report_ws["A14"] = "TOP 5 BY FOLLOWERS"

for row in report_ws["A14:F14"]:
    for cell in row:
        apply_fill(cell, INK_2)

report_ws["A14"].font = Font(
    name="Aptos",
    size=10,
    bold=True,
    color=WHITE,
)


headers = ["Rank", "Account", "Followers"]

for col, header in enumerate(headers, 1):
    cell = report_ws.cell(15, col)
    cell.value = header
    apply_fill(cell, "EEF2F6")

    cell.font = Font(
        name="Aptos",
        size=9,
        bold=True,
        color=MUTED,
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


top_followers = report[:5]

for index, row in enumerate(top_followers, 16):
    values = [
        row["_rank"],
        f"@{row.get('handle', '')}",
        row["_followers"],
    ]

    for col, value in enumerate(values, 1):
        cell = report_ws.cell(index, col)

        cell.value = value

        apply_fill(cell, CARD)

        cell.font = Font(
            name="Aptos",
            size=10,
            color=TEXT,
            bold=(
                row.get("handle") == NBA_HANDLE
            ),
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = Border(
            bottom=Side(
                style="thin",
                color=LINE,
            )
        )

    report_ws.cell(index, 3).number_format = "#,##0"


# ============================================================
# FASTEST GROWTH
# ============================================================

report_ws.merge_cells("G14:L14")

report_ws["G14"] = "FASTEST 30-DAY+ GROWTH"

for row in report_ws["G14:L14"]:
    for cell in row:
        apply_fill(cell, INK_2)

report_ws["G14"].font = Font(
    name="Aptos",
    size=10,
    bold=True,
    color=WHITE,
)


growth_headers = [
    "Account",
    "Change",
    "Growth",
    "Period",
]

for col, header in enumerate(growth_headers, 7):
    cell = report_ws.cell(15, col)
    cell.value = header

    apply_fill(cell, "EEF2F6")

    cell.font = Font(
        name="Aptos",
        size=9,
        bold=True,
        color=MUTED,
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )


growth_rows = sorted(
    [
        row
        for row in report
        if row["_growth"] is not None
    ],
    key=lambda row: row["_growth"],
    reverse=True,
)[:5]


for index, row in enumerate(growth_rows, 16):
    values = [
        f"@{row.get('handle', '')}",
        row["_change"],
        row["_growth"],
        (
            f"{row['_comparison_days']}d"
            if row["_comparison_days"]
            else ""
        ),
    ]

    for col, value in enumerate(values, 7):
        cell = report_ws.cell(index, col)

        cell.value = value

        apply_fill(cell, CARD)

        cell.font = Font(
            name="Aptos",
            size=10,
            color=TEXT,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        cell.border = Border(
            bottom=Side(
                style="thin",
                color=LINE,
            )
        )

    report_ws.cell(index, 8).number_format = (
        "+#,##0;-#,##0;0"
    )

    report_ws.cell(index, 9).number_format = (
        "+0.00%;-0.00%;0.00%"
    )


# ============================================================
# ISSUE BLOCK
# ============================================================

report_ws.merge_cells("A23:L23")

report_ws["A23"] = "CURRENT PROFILE ISSUES"

for row in report_ws["A23:L23"]:
    for cell in row:
        apply_fill(cell, AMBER_BG)

report_ws["A23"].font = Font(
    name="Aptos",
    size=9,
    bold=True,
    color=AMBER_TEXT,
)


report_ws.merge_cells("A24:L25")

issue_cell = report_ws["A24"]

if errors:
    handles = [
        str(row.get("handle", "")).strip()
        for row in errors
        if row.get("handle")
    ]

    handles = [
        f"@{handle}"
        for handle in handles
    ]

    issue_cell.value = (
        f"{', '.join(handles)} — profile data "
        f"could not be fully collected. "
        f"The report still includes "
        f"{len(report)} valid profiles and "
        f"failed accounts will continue to be "
        f"checked automatically each week."
    )
else:
    issue_cell.value = (
        "No profile collection issues were "
        "reported in the latest run."
    )


for row in report_ws["A24:L25"]:
    for cell in row:
        apply_fill(cell, AMBER_LIGHT)

issue_cell.font = Font(
    name="Aptos",
    size=10,
    color=AMBER_TEXT,
)

issue_cell.alignment = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=True,
)


# ============================================================
# COLUMN WIDTHS
# ============================================================

widths = {
    "A": 10,
    "B": 23,
    "C": 14,
    "D": 12,
    "E": 12,
    "F": 12,
    "G": 24,
    "H": 12,
    "I": 12,
    "J": 12,
    "K": 12,
    "L": 12,
}

for col, width in widths.items():
    report_ws.column_dimensions[col].width = width


for row, height in {
    1: 22,
    2: 22,
    3: 22,
    4: 20,
    7: 24,
    8: 30,
    9: 30,
    10: 30,
    11: 22,
    14: 24,
    15: 22,
    23: 22,
    24: 30,
    25: 30,
}.items():
    report_ws.row_dimensions[row].height = height


# ============================================================
# COMPETITOR RANKING SHEET
# ============================================================

ranking_headers = [
    "Rank",
    "Handle",
    "Followers",
    "30-Day+ Change",
    "Growth %",
    "Comparison Date",
    "Period",
    "Total Posts",
    "Posts Since Previous Snapshot",
    "Ahead of NBA",
]


ranking_ws.append(ranking_headers)

for row in report:
    ranking_ws.append(
        [
            row["_rank"],
            row.get("handle"),
            row["_followers"],
            row["_change"],
            row["_growth"],
            row.get("comparison_date"),
            row["_comparison_days"],
            row["_posts"],
            row["_posts_change"],
            row.get("ahead_of_nba"),
        ]
    )


style_header_row(ranking_ws)
style_body(ranking_ws)

ranking_ws.freeze_panes = "A2"
ranking_ws.sheet_view.showGridLines = False


for row in ranking_ws.iter_rows(
    min_row=2,
    max_row=ranking_ws.max_row,
):
    row[2].number_format = "#,##0"
    row[3].number_format = "+#,##0;-#,##0;0"
    row[4].number_format = "+0.00%;-0.00%;0.00%"
    row[7].number_format = "#,##0"
    row[8].number_format = "+#,##0;-#,##0;0"

    growth_cell = row[4]

    if growth_cell.value is not None:
        growth_cell.font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=(
                GREEN
                if growth_cell.value >= 0
                else RED
            ),
        )


ranking_widths = {
    "A": 8,
    "B": 28,
    "C": 14,
    "D": 16,
    "E": 12,
    "F": 16,
    "G": 10,
    "H": 14,
    "I": 28,
    "J": 13,
}

for col, width in ranking_widths.items():
    ranking_ws.column_dimensions[col].width = width


# ============================================================
# GROWTH SHEET
# ============================================================

growth_ws.append(
    [
        "Handle",
        "Current Followers",
        "Comparison Followers",
        "Change",
        "Growth %",
        "Comparison Date",
        "Period",
        "Current Rank",
    ]
)


all_growth_rows = sorted(
    report,
    key=lambda row: (
        row["_growth"]
        if row["_growth"] is not None
        else -999
    ),
    reverse=True,
)


for row in all_growth_rows:
    growth_ws.append(
        [
            row.get("handle"),
            row["_followers"],
            row["_comparison_followers"],
            row["_change"],
            row["_growth"],
            row.get("comparison_date"),
            row["_comparison_days"],
            row["_rank"],
        ]
    )


style_header_row(growth_ws)
style_body(growth_ws)

growth_ws.freeze_panes = "A2"
growth_ws.sheet_view.showGridLines = False


for row in growth_ws.iter_rows(
    min_row=2,
    max_row=growth_ws.max_row,
):
    row[1].number_format = "#,##0"
    row[2].number_format = "#,##0"
    row[3].number_format = "+#,##0;-#,##0;0"
    row[4].number_format = "+0.00%;-0.00%;0.00%"

    if row[4].value is not None:
        row[4].font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=(
                GREEN
                if row[4].value >= 0
                else RED
            ),
        )


growth_widths = {
    "A": 28,
    "B": 17,
    "C": 19,
    "D": 13,
    "E": 12,
    "F": 16,
    "G": 10,
    "H": 12,
}

for col, width in growth_widths.items():
    growth_ws.column_dimensions[col].width = width


# ============================================================
# POSTING ACTIVITY
# ============================================================

activity_ws.append(
    [
        "Handle",
        "Current Total Posts",
        "Posts Since Previous Snapshot",
        "Previous Snapshot Date",
    ]
)


for row in report:
    activity_ws.append(
        [
            row.get("handle"),
            row["_posts"],
            row["_posts_change"],
            row.get("previous_post_snapshot_date"),
        ]
    )


style_header_row(activity_ws)
style_body(activity_ws)

activity_ws.freeze_panes = "A2"
activity_ws.sheet_view.showGridLines = False


for row in activity_ws.iter_rows(
    min_row=2,
    max_row=activity_ws.max_row,
):
    row[1].number_format = "#,##0"
    row[2].number_format = "+#,##0;-#,##0;0"


activity_widths = {
    "A": 28,
    "B": 20,
    "C": 29,
    "D": 23,
}

for col, width in activity_widths.items():
    activity_ws.column_dimensions[col].width = width


# ============================================================
# ISSUES
# ============================================================

issue_headers = [
    "Date Checked",
    "Handle",
    "Status",
    "Reason",
]


issues_ws.append(issue_headers)


if errors:
    for error in errors:
        date_checked = (
            error.get("date_checked")
            or latest_date
        )

        handle = error.get("handle", "")

        reason = (
            error.get("error")
            or error.get("reason")
            or error.get("message")
            or error.get("status")
            or "Profile collection failed"
        )

        issues_ws.append(
            [
                date_checked,
                handle,
                "Failed",
                reason,
            ]
        )

else:
    issues_ws.append(
        [
            latest_date,
            "",
            "OK",
            "No profile collection issues.",
        ]
    )


style_header_row(issues_ws)
style_body(issues_ws)

issues_ws.sheet_view.showGridLines = False

issues_ws.column_dimensions["A"].width = 16
issues_ws.column_dimensions["B"].width = 28
issues_ws.column_dimensions["C"].width = 12
issues_ws.column_dimensions["D"].width = 60


for row in issues_ws.iter_rows(
    min_row=2,
    max_row=issues_ws.max_row,
):
    if row[2].value == "Failed":
        for cell in row:
            apply_fill(cell, AMBER_LIGHT)

        row[2].font = Font(
            name="Aptos",
            size=10,
            bold=True,
            color=AMBER_TEXT,
        )


# ============================================================
# HISTORY
# ============================================================

history_ws.append(
    [
        "Date",
        "Handle",
        "Followers",
        "Source",
    ]
)


history_rows = []


for row in historical:
    handle = (
        row.get("handle")
        or row.get("Handle")
        or row.get("account")
        or row.get("Account")
    )

    followers = (
        row.get("followers")
        or row.get("Followers")
    )

    date = (
        row.get("date")
        or row.get("Date")
        or row.get("snapshot_date")
    )

    if handle and followers:
        history_rows.append(
            [
                date,
                handle,
                as_int(followers),
                "historical",
            ]
        )


for row in snapshots:
    handle = row.get("handle")
    followers = as_int(row.get("followers"))
    date = row.get("date")

    if handle and followers is not None:
        history_rows.append(
            [
                date,
                handle,
                followers,
                row.get("source") or "automated",
            ]
        )


history_rows.sort(
    key=lambda row: (
        str(row[0]),
        str(row[1]),
    )
)


for row in history_rows:
    history_ws.append(row)


style_header_row(history_ws)
style_body(history_ws)

history_ws.freeze_panes = "A2"
history_ws.sheet_view.showGridLines = False

history_ws.column_dimensions["A"].width = 16
history_ws.column_dimensions["B"].width = 28
history_ws.column_dimensions["C"].width = 14
history_ws.column_dimensions["D"].width = 24


for row in history_ws.iter_rows(
    min_row=2,
    max_row=history_ws.max_row,
):
    row[2].number_format = "#,##0"


# ============================================================
# CHART DATA
# ============================================================

chart_start = 47

report_ws.cell(
    chart_start,
    1,
    "Handle",
)

report_ws.cell(
    chart_start,
    2,
    "Followers",
)


for index, row in enumerate(
    top_followers,
    chart_start + 1,
):
    report_ws.cell(
        index,
        1,
        f"@{row.get('handle')}",
    )

    report_ws.cell(
        index,
        2,
        row["_followers"],
    )


growth_chart_col = 7

report_ws.cell(
    chart_start,
    growth_chart_col,
    "Handle",
)

report_ws.cell(
    chart_start,
    growth_chart_col + 1,
    "Growth",
)


for index, row in enumerate(
    growth_rows,
    chart_start + 1,
):
    report_ws.cell(
        index,
        growth_chart_col,
        f"@{row.get('handle')}",
    )

    report_ws.cell(
        index,
        growth_chart_col + 1,
        row["_growth"],
    )

    report_ws.cell(
        index,
        growth_chart_col + 1,
    ).number_format = "0.00%"


# ============================================================
# FOLLOWERS CHART
# ============================================================

followers_chart = BarChart()

followers_chart.type = "bar"
followers_chart.style = 10

followers_chart.title = "Follower Landscape"
followers_chart.y_axis.title = ""
followers_chart.x_axis.title = ""

followers_chart.height = 7.2
followers_chart.width = 13.5

followers_chart.legend = None

followers_chart.varyColors = False


data = Reference(
    report_ws,
    min_col=2,
    min_row=chart_start,
    max_row=chart_start + len(top_followers),
)

cats = Reference(
    report_ws,
    min_col=1,
    min_row=chart_start + 1,
    max_row=chart_start + len(top_followers),
)

followers_chart.add_data(
    data,
    titles_from_data=True,
)

followers_chart.set_categories(cats)

followers_chart.dLbls = DataLabelList()
followers_chart.dLbls.showVal = True


if followers_chart.series:
    followers_chart.series[0].graphicalProperties.solidFill = CYAN
    followers_chart.series[0].graphicalProperties.line.solidFill = CYAN


report_ws.add_chart(
    followers_chart,
    "A28",
)


# ============================================================
# GROWTH CHART
# ============================================================

growth_chart = BarChart()

growth_chart.type = "bar"
growth_chart.style = 10

growth_chart.title = "Fastest 30-Day+ Growth"
growth_chart.y_axis.title = ""
growth_chart.x_axis.title = ""

growth_chart.height = 7.2
growth_chart.width = 13.5

growth_chart.legend = None


data = Reference(
    report_ws,
    min_col=growth_chart_col + 1,
    min_row=chart_start,
    max_row=chart_start + len(growth_rows),
)

cats = Reference(
    report_ws,
    min_col=growth_chart_col,
    min_row=chart_start + 1,
    max_row=chart_start + len(growth_rows),
)


growth_chart.add_data(
    data,
    titles_from_data=True,
)

growth_chart.set_categories(cats)

growth_chart.dLbls = DataLabelList()
growth_chart.dLbls.showVal = True


if growth_chart.series:
    growth_chart.series[0].graphicalProperties.solidFill = GREEN
    growth_chart.series[0].graphicalProperties.line.solidFill = GREEN


report_ws.add_chart(
    growth_chart,
    "G28",
)


# Hide chart source rows.
for row in range(
    chart_start,
    chart_start + 7,
):
    report_ws.row_dimensions[row].hidden = True


# ============================================================
# PRINT / VIEW
# ============================================================

report_ws.freeze_panes = "A7"

report_ws.page_setup.orientation = "landscape"
report_ws.page_setup.fitToWidth = 1
report_ws.page_setup.fitToHeight = 0

report_ws.sheet_properties.pageSetUpPr.fitToPage = True

report_ws.print_area = "A1:L43"

report_ws.sheet_view.zoomScale = 90


# ============================================================
# SAVE
# ============================================================

OUTPUT_XLSX.parent.mkdir(
    parents=True,
    exist_ok=True,
)

wb.save(OUTPUT_XLSX)

print("=" * 78)
print("NBA MODERN EXCEL REPORT")
print("=" * 78)
print()
print(f"Report date:     {latest_date}")
print(f"NBA followers:   {nba_followers:,}")
print(f"NBA rank:        {nba_rank} of {len(report)}")
print(f"Accounts ahead:  {accounts_ahead}")
print()
print(f"Saved:")
print(OUTPUT_XLSX)
print()
print("=" * 78)

